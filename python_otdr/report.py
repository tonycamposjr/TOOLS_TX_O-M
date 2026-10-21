import os
import json
from typing import Dict, Any

class OTDRReportGenerator:
    """
    Gerador de relatórios de certificação OTDR.
    Gera relatórios em PDF, Excel (.xlsx) e JSON estruturado.
    """

    def __init__(self, sor_data: Dict[str, Any], validation_data: Dict[str, Any]):
        self.sor_data = sor_data
        self.validation_data = validation_data
        self.gen = sor_data.get("general_params", {})
        self.fxd = sor_data.get("fixed_params", {})
        self.link = sor_data.get("link_params", {})
        self.trace = sor_data.get("trace_data", {})
        self.events = validation_data.get("evaluated_events", [])

    def _get_base_filename(self) -> str:
        """Gera nome de arquivo amigável baseado no ID da fibra e data."""
        cable = self.gen.get("cable_id", "CABO").replace(" ", "_")
        fiber = self.gen.get("fiber_id", "FIBRA").replace(" ", "_")
        dt_str = self.fxd.get("date_time", "").replace(" ", "_").replace(":", "-")[:10]
        if not dt_str:
            dt_str = "RELATORIO"
        return f"CERTIFICACAO_OTDR_{cable}_{fiber}_{dt_str}"

    def generate_plot_image(self, output_png_path: str):
        """Gera o gráfico do traço OTDR (distância x potência) via Matplotlib."""
        try:
            import matplotlib
            matplotlib.use('Agg')
            import matplotlib.pyplot as plt

            distances = self.trace.get("distances_km", [])
            powers = self.trace.get("data_points_db", [])

            fig, ax = plt.subplots(figsize=(10, 4.5), dpi=150)
            fig.patch.set_facecolor('#f8fafc')
            ax.set_facecolor('#ffffff')

            if distances and powers:
                ax.plot(distances, powers, color='#2563eb', linewidth=1.5, label='Traço OTDR (dB)')

            # Adicionar marcadores de eventos
            for ev in self.events:
                d = ev.get("distance_km", 0)
                ev_num = ev.get("event_number", 1)
                st = ev.get("status", "APROVADO")
                color = '#16a34a' if st == "APROVADO" else '#dc2626'
                
                # Linha vertical discreta
                ax.axvline(x=d, color=color, linestyle='--', alpha=0.6, linewidth=1)
                # Anotação com o número do evento
                ax.annotate(
                    f"E{ev_num}",
                    xy=(d, ax.get_ylim()[1] if len(powers) else 30),
                    xytext=(d, ax.get_ylim()[1] - 3 if len(powers) else 27),
                    fontsize=8,
                    fontweight='bold',
                    color=color,
                    ha='center',
                    bbox=dict(boxstyle="circle,pad=0.2", fc="#ffffff", ec=color, lw=1)
                )

            ax.set_title(f"Traço OTDR - Fibra {self.gen.get('fiber_id')} ({self.fxd.get('wavelength_nm')} nm)", fontsize=12, fontweight='bold', pad=12)
            ax.set_xlabel("Distância (km)", fontsize=10)
            ax.set_ylabel("Nível de Potência (dB)", fontsize=10)
            ax.grid(True, linestyle=':', alpha=0.6)
            ax.legend(loc='upper right', fontsize=9)

            plt.tight_layout()
            plt.savefig(output_png_path, dpi=150)
            plt.close()
            return output_png_path
        except Exception as e:
            print(f"Aviso: Não foi possível gerar gráfico matplotlib: {e}")
            return None

    def export_json(self, output_dir: str) -> str:
        """Exporta os dados completos e resultado da validação para arquivo JSON."""
        os.makedirs(output_dir, exist_ok=True)
        filename = f"{self._get_base_filename()}.json"
        filepath = os.path.join(output_dir, filename)

        full_export = {
            "sor_data": self.sor_data,
            "certification": self.validation_data
        }

        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(full_export, f, indent=2, ensure_ascii=False)

        return filepath

    def export_excel(self, output_dir: str) -> str:
        """Exporta os dados formatados e tabela de eventos para Excel (.xlsx)."""
        os.makedirs(output_dir, exist_ok=True)
        filename = f"{self._get_base_filename()}.xlsx"
        filepath = os.path.join(output_dir, filename)

        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Relatório OTDR"

            # Estilos
            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            pass_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
            fail_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            bold_font = Font(bold=True)

            # Cabeçalho do Documento
            ws["A1"] = "RELATÓRIO DE CERTIFICAÇÃO DE FIBRA ÓPTICA (OTDR)"
            ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")

            ws["A3"] = "ID do Cabo:"
            ws["B3"] = self.gen.get("cable_id")
            ws["A4"] = "ID da Fibra:"
            ws["B4"] = self.gen.get("fiber_id")
            ws["A5"] = "Comprimento de Onda:"
            ws["B5"] = f"{self.fxd.get('wavelength_nm')} nm"
            ws["A6"] = "Data do Teste:"
            ws["B6"] = self.fxd.get("date_time")

            ws["D3"] = "Técnico:"
            ws["E3"] = self.gen.get("operator")
            ws["D4"] = "Distância Total:"
            ws["E4"] = f"{self.link.get('total_distance_km')} km"
            ws["D5"] = "Perda Total:"
            ws["E5"] = f"{self.link.get('total_loss_db')} dB"
            ws["D6"] = "STATUS GERAL:"
            
            st_cell = ws["E6"]
            st_cell.value = self.validation_data.get("overall_status")
            st_cell.font = Font(bold=True, color="047857" if st_cell.value == "APROVADO" else "B91C1C")
            st_cell.fill = pass_fill if st_cell.value == "APROVADO" else fail_fill

            # Tabela de Eventos
            headers = ["# Evento", "Distância (km)", "Tipo de Evento", "Perda (dB)", "Refletância (dB)", "Status"]
            row_idx = 9
            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            for ev in self.events:
                row_idx += 1
                ws.cell(row=row_idx, column=1, value=ev.get("event_number"))
                ws.cell(row=row_idx, column=2, value=ev.get("distance_km"))
                ws.cell(row=row_idx, column=3, value=ev.get("event_type"))
                ws.cell(row=row_idx, column=4, value=ev.get("loss_db"))
                ws.cell(row=row_idx, column=5, value=str(ev.get("reflectance_db")))
                
                st = ws.cell(row=row_idx, column=6, value=ev.get("status"))
                st.alignment = Alignment(horizontal="center")
                st.fill = pass_fill if ev.get("status") == "APROVADO" else fail_fill

            wb.save(filepath)
            return filepath
        except Exception as e:
            print(f"Aviso ao exportar Excel via openpyxl: {e}")
            return self.export_json(output_dir)

    def export_pdf(self, output_dir: str) -> str:
        """Gera o relatório em PDF estruturado usando ReportLab."""
        os.makedirs(output_dir, exist_ok=True)
        filename = f"{self._get_base_filename()}.pdf"
        filepath = os.path.join(output_dir, filename)

        plot_path = os.path.join(output_dir, "temp_otdr_trace.png")
        self.generate_plot_image(plot_path)

        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib import colors

            doc = SimpleDocTemplate(filepath, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
            story = []
            styles = getSampleStyleSheet()

            # Estilo personalizado
            title_style = ParagraphStyle(
                'DocTitle',
                parent=styles['Heading1'],
                fontSize=18,
                textColor=colors.HexColor('#1e3a8a'),
                spaceAfter=12
            )

            story.append(Paragraph("RELATÓRIO DE CERTIFICAÇÃO DE FIBRA ÓPTICA (OTDR)", title_style))
            story.append(Spacer(1, 10))

            # Status e dados gerais
            st_val = self.validation_data.get("overall_status")
            st_color = colors.HexColor("#16a34a") if st_val == "APROVADO" else colors.HexColor("#dc2626")

            info_data = [
                [
                    Paragraph(f"<b>Cabo:</b> {self.gen.get('cable_id')}", styles['Normal']),
                    Paragraph(f"<b>Técnico:</b> {self.gen.get('operator')}", styles['Normal'])
                ],
                [
                    Paragraph(f"<b>Fibra:</b> {self.gen.get('fiber_id')}", styles['Normal']),
                    Paragraph(f"<b>Comprimento de Onda:</b> {self.fxd.get('wavelength_nm')} nm", styles['Normal'])
                ],
                [
                    Paragraph(f"<b>Distância Total:</b> {self.link.get('total_distance_km')} km", styles['Normal']),
                    Paragraph(f"<b>Perda Total:</b> {self.link.get('total_loss_db')} dB", styles['Normal'])
                ],
                [
                    Paragraph(f"<b>Data:</b> {self.fxd.get('date_time')}", styles['Normal']),
                    Paragraph(f"<b>STATUS:</b> <font color='{st_color.hexval()}'><b>{st_val}</b></font>", styles['Normal'])
                ]
            ]

            info_table = Table(info_data, colWidths=[270, 270])
            info_table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f1f5f9')),
                ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
                ('PADDING', (0,0), (-1,-1), 8),
            ]))
            story.append(info_table)
            story.append(Spacer(1, 15))

            # Gráfico do Traço
            if os.path.exists(plot_path):
                story.append(Image(plot_path, width=520, height=230))
                story.append(Spacer(1, 15))

            # Tabela de Eventos
            story.append(Paragraph("<b>Lista de Eventos Identificados:</b>", styles['Heading3']))
            story.append(Spacer(1, 6))

            evt_data = [["#", "Distância (km)", "Tipo de Evento", "Perda (dB)", "Refletância", "Status"]]
            for ev in self.events:
                evt_data.append([
                    str(ev.get("event_number")),
                    f"{ev.get('distance_km'):.3f}",
                    ev.get("event_type"),
                    f"{ev.get('loss_db'):.3f}",
                    str(ev.get("reflectance_db")),
                    ev.get("status")
                ])

            evt_table = Table(evt_data, colWidths=[30, 95, 150, 80, 85, 80])
            evt_table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e3a8a')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.white),
                ('ALIGN', (0,0), (-1,-1), 'CENTER'),
                ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                ('BOTTOMPADDING', (0,0), (-1,0), 6),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
            ]))
            story.append(evt_table)

            doc.build(story)

            if os.path.exists(plot_path):
                os.remove(plot_path)

            return filepath
        except Exception as e:
            print(f"Erro ao gerar PDF com ReportLab: {e}. Gerando saída estruturada alternativamente.")
            return self.export_json(output_dir)
